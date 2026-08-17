import io

import pytest
from docx import Document

from src.core.exceptions import UnsupportedFileTypeError
from src.services.document_parser import DOCX_CONTENT_TYPE, extract_text


def test_extract_text_raises_for_unsupported_type():
    with pytest.raises(UnsupportedFileTypeError):
        extract_text(b"hello", "text/plain")


def test_extract_text_reads_docx_paragraphs():
    buffer = io.BytesIO()
    doc = Document()
    doc.add_paragraph("First paragraph.")
    doc.add_paragraph("Second paragraph.")
    doc.save(buffer)

    text = extract_text(buffer.getvalue(), DOCX_CONTENT_TYPE)

    assert "First paragraph." in text
    assert "Second paragraph." in text
